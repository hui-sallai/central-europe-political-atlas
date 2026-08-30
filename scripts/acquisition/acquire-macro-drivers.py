"""Build the v1.51 macro-driver research layer from official bulk/API sources.

The script deliberately keeps observed movements separate from identified shocks.
It stores latest-revised monthly data from 2015 onward and never interpolates.
"""

from __future__ import annotations

import csv
import hashlib
import json
import math
import os
import re
import urllib.parse
import urllib.request
import zipfile
from collections import defaultdict
from datetime import date
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
CACHE = ROOT / ".tmp-macro-drivers"
OUT = ROOT / "src" / "data" / "macro-drivers"
START = "2015-01"
TODAY = date.today().isoformat()
SCHEMA_VERSION = "v1.51"

COUNTRIES = {
    "AT": "austria",
    "CZ": "czechia",
    "DE": "germany",
    "HR": "croatia",
    "HU": "hungary",
    "PL": "poland",
    "RO": "romania",
    "RS": "serbia",
    "SI": "slovenia",
    "SK": "slovakia",
}
EURO_AREA = {"austria", "croatia", "germany", "slovakia", "slovenia"}
NON_EURO = {"czechia", "hungary", "poland", "romania", "serbia"}
CURRENCIES = {"czechia": "CZK", "hungary": "HUF", "poland": "PLN", "romania": "RON", "serbia": "RSD"}

BIS_FILES = {
    "cbpol": "https://data.bis.org/static/bulk/WS_CBPOL_csv_col.zip",
    "xru": "https://data.bis.org/static/bulk/WS_XRU_csv_col.zip",
    "eer": "https://data.bis.org/static/bulk/WS_EER_csv_col.zip",
}
PINK_SHEET_URL = "https://thedocs.worldbank.org/en/doc/74e8be41ceb20fa0da750cda2f6b9e4e-0050012026/related/CMO-Historical-Data-Monthly.xlsx"


def write_json(name: str, payload: object) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / name).write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def download(url: str, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    if not path.exists():
        request = urllib.request.Request(url, headers={"User-Agent": "Central-Europe-Political-Atlas/1.5 research acquisition"})
        with urllib.request.urlopen(request, timeout=180) as response, path.open("wb") as output:
            output.write(response.read())
    return path


def unzip_csv(zip_path: Path, folder: Path) -> Path:
    folder.mkdir(parents=True, exist_ok=True)
    candidates = list(folder.glob("*.csv"))
    if not candidates:
        with zipfile.ZipFile(zip_path) as archive:
            archive.extractall(folder)
        candidates = list(folder.rglob("*.csv"))
    if not candidates:
        raise RuntimeError(f"No CSV found in {zip_path}")
    return candidates[0]


def monthly_fields(row: dict[str, str]) -> dict[str, float | None]:
    values: dict[str, float | None] = {}
    for key, raw in row.items():
        if not re.fullmatch(r"\d{4}-\d{2}", key or "") or key < START:
            continue
        try:
            values[key] = float(raw) if raw not in (None, "") else None
        except ValueError:
            values[key] = None
    return values


def month_sequence(start: str, end: str) -> list[str]:
    year, month = map(int, start.split("-"))
    end_year, end_month = map(int, end.split("-"))
    result = []
    while (year, month) <= (end_year, end_month):
        result.append(f"{year:04d}-{month:02d}")
        month += 1
        if month == 13:
            month = 1
            year += 1
    return result


def month_ordinal(period: str) -> int:
    year, month = map(int, period.split("-"))
    return year * 12 + month - 1


def month_distance(previous_period: str, current_period: str) -> int:
    return month_ordinal(current_period) - month_ordinal(previous_period)


def shift_month(period: str, delta: int) -> str:
    ordinal = month_ordinal(period) + delta
    year, zero_based_month = divmod(ordinal, 12)
    return f"{year:04d}-{zero_based_month + 1:02d}"


def is_consecutive_month(previous_period: str, current_period: str) -> bool:
    return month_distance(previous_period, current_period) == 1


def eurostat(dataset: str, params: dict[str, str], geos: list[str]) -> tuple[list[dict], dict, str, Path]:
    query: list[tuple[str, str]] = [("format", "JSON"), ("lang", "en"), ("sinceTimePeriod", START)]
    query.extend(params.items())
    query.extend(("geo", geo) for geo in geos)
    url = f"https://ec.europa.eu/eurostat/api/dissemination/statistics/1.0/data/{dataset}?{urllib.parse.urlencode(query)}"
    cache = CACHE / f"eurostat-{dataset}-{'-'.join(params.values())}.json"
    download(url, cache)
    payload = json.loads(cache.read_text(encoding="utf-8"))
    ids = payload["id"]
    sizes = payload["size"]
    time_id = ids[-1]
    dimensions = []
    for dim_id in ids[:-1]:
        index = payload["dimension"][dim_id]["category"]["index"]
        dimensions.append((dim_id, [key for key, _ in sorted(index.items(), key=lambda item: item[1])]))
    time_index = payload["dimension"][time_id]["category"]["index"]
    times = [key for key, _ in sorted(time_index.items(), key=lambda item: item[1])]
    strides = [1] * len(sizes)
    for index in range(len(sizes) - 2, -1, -1):
        strides[index] = strides[index + 1] * sizes[index + 1]
    rows: list[dict] = []

    def walk(dim_index: int, combo: dict[str, str], offset: int) -> None:
        if dim_index == len(dimensions):
            for time_position, period in enumerate(times):
                flat = offset + time_position * strides[-1]
                value = payload.get("value", {}).get(str(flat), payload.get("value", {}).get(flat))
                rows.append({**combo, "period": period, "value": value})
            return
        dim_id, keys = dimensions[dim_index]
        for position, key in enumerate(keys):
            walk(dim_index + 1, {**combo, dim_id: key}, offset + position * strides[dim_index])

    walk(0, {}, 0)
    return rows, payload, url, cache


def observation(driver_id: str, country: str | None, scope: str, period: str, value: float | None,
                unit: str, transformation: str, economic_role: str, source: str, source_url: str,
                source_dataset: str, definition_version: str, aggregation_method: str,
                orientation: str, data_status: str = "official", **extra: object) -> dict:
    area = country or scope
    numeric_value = None if value is None or not math.isfinite(float(value)) else round(float(value), 8)
    return {
        "observation_id": f"md:{driver_id}:{transformation}:{area}:{period}",
        "driver_id": driver_id,
        "country": country,
        "scope": scope,
        "frequency": "monthly",
        "period": period,
        "value": numeric_value,
        "unit": unit,
        "transformation": transformation,
        "source": source,
        "source_url": source_url,
        "source_dataset": source_dataset,
        "source_reliability": "A",
        "definition_version": definition_version,
        "revision_status": "latest_revised",
        "vintage_available": False,
        "data_status": data_status if numeric_value is not None or data_status == "computed" else "pending",
        "updated_at": TODAY,
        "economic_role": economic_role,
        "role": economic_role,
        "identification_status": extra.pop("identification_status", "observed_driver"),
        "series_instance_id": extra.get("series_instance_id") or extra.get("source_series_instance_id") or extra.get("source_series_id") or f"{driver_id}:{area}",
        "applicability_scope": extra.get("applicability_scope") or scope,
        "applicable_country_ids": extra.get("applicable_country_ids") or ([country] if country else []),
        "shared_series": bool(extra.get("shared_series", country is None)),
        "independent_cross_section_unit": bool(extra.get("independent_cross_section_unit", country is not None)),
        "aggregation_method": aggregation_method,
        "orientation": orientation,
        **extra,
    }


def add_changes(records: list[dict], driver_id: str, country: str | None, scope: str,
                values: dict[str, float | None], unit: str, economic_role: str, source: str,
                source_url: str, source_dataset: str, definition: str, aggregation: str,
                orientation: str, include_yoy: bool = False, **extra: object) -> None:
    periods = sorted(values)
    derived_extra = {key: value for key, value in extra.items() if key != "calculation_method"}
    if extra.get("calculation_method"):
        derived_extra["source_level_calculation_method"] = extra["calculation_method"]
    for index, period in enumerate(periods):
        value = values[period]
        level = observation(driver_id, country, scope, period, value, unit, "level", economic_role, source,
                            source_url, source_dataset, definition, aggregation, orientation, **extra)
        records.append(level)
        previous_period = shift_month(period, -1)
        previous = values.get(previous_period)
        available_prior_periods = [candidate for candidate in periods if candidate < period and values.get(candidate) is not None]
        actual_previous_period = available_prior_periods[-1] if available_prior_periods else None
        trace_previous_period = previous_period if previous is not None else actual_previous_period
        previous_id = f"md:{driver_id}:level:{country or scope}:{trace_previous_period}" if trace_previous_period else None
        if period == periods[0] and previous_period not in values:
            alignment, derivation, reason = "warmup", "warmup", "transformation_warmup"
        elif previous_period not in values or previous is None or value is None:
            alignment, derivation, reason = "source_gap", "gap_blocked", "non_consecutive_source_period"
        elif value <= 0 or previous <= 0:
            alignment, derivation, reason = "definition_break", "definition_blocked", "log_change_requires_positive_values"
        else:
            alignment, derivation, reason = "aligned", "valid", None
        change = 100 * math.log(value / previous) if derivation == "valid" else None
        records.append(observation(driver_id, country, scope, period, change, "%", "monthly_log_change",
                                   economic_role,
                                   source, source_url, source_dataset, definition, aggregation,
                                   "positive means an increase in the source level", data_status="computed",
                                   calculation_method="100 * ln(value_t / value_t-1)",
                                   identification_status="shock_candidate" if economic_role == "external_common_driver" else "observed_driver",
                                   source_periods=[trace_previous_period, period] if trace_previous_period else [period],
                                   source_observation_ids=[previous_id, level["observation_id"]] if previous_id else [level["observation_id"]],
                                   required_lag_months=1,
                                   actual_lag_months=month_distance(trace_previous_period, period) if trace_previous_period else None,
                                   temporal_alignment_status=alignment,
                                   derivation_status=derivation,
                                   availability_reason=reason,
                                   **derived_extra))
        if include_yoy:
            previous_12_period = shift_month(period, -12)
            previous_12 = values.get(previous_12_period)
            previous_12_id = f"md:{driver_id}:level:{country or scope}:{previous_12_period}"
            source_has_started = previous_12_period >= periods[0]
            if not source_has_started:
                alignment_12, derivation_12, reason_12 = "warmup", "warmup", "transformation_warmup"
            elif previous_12_period not in values or previous_12 is None or value is None:
                alignment_12, derivation_12, reason_12 = "source_gap", "gap_blocked", "missing_exact_12_month_source_period"
            elif value <= 0 or previous_12 <= 0:
                alignment_12, derivation_12, reason_12 = "definition_break", "definition_blocked", "log_change_requires_positive_values"
            else:
                alignment_12, derivation_12, reason_12 = "aligned", "valid", None
            yoy = 100 * math.log(value / previous_12) if derivation_12 == "valid" else None
            records.append(observation(driver_id, country, scope, period, yoy, "%", "12m_log_change",
                                       economic_role,
                                       source, source_url, source_dataset, definition, aggregation,
                                       "positive means an increase in the source level", data_status="computed",
                                       calculation_method="100 * ln(value_t / value_t-12)",
                                       identification_status="observed_driver",
                                       source_periods=[previous_12_period, period],
                                       source_observation_ids=[previous_12_id, level["observation_id"]],
                                       required_lag_months=12,
                                       actual_lag_months=12 if previous_12_period in values else None,
                                       temporal_alignment_status=alignment_12,
                                       derivation_status=derivation_12,
                                       availability_reason=reason_12,
                                       **derived_extra))


def build() -> None:
    CACHE.mkdir(parents=True, exist_ok=True)
    OUT.mkdir(parents=True, exist_ok=True)
    observations: list[dict] = []
    manifests: dict[str, dict] = {}

    bis_paths: dict[str, tuple[Path, Path]] = {}
    for key, url in BIS_FILES.items():
        archive = download(url, CACHE / f"{key}.zip")
        source = unzip_csv(archive, CACHE / key)
        bis_paths[key] = (archive, source)

    # Policy rates: resolved country views retain one canonical shared-series identity for ECB rows.
    with bis_paths["cbpol"][1].open(encoding="utf-8-sig", newline="") as handle:
        policy_rows = {row["REF_AREA"]: row for row in csv.DictReader(handle) if row["FREQ"] == "M"}
    policy_series_manifest = []
    euro_row = policy_rows["XM"]
    for geo, country in COUNTRIES.items():
        use_common = country in EURO_AREA
        source_row = euro_row if use_common else policy_rows.get(geo)
        if source_row is None:
            policy_series_manifest.append({"country": country, "status": "missing", "reason": "No comparable BIS monthly policy series"})
            continue
        values = monthly_fields(source_row)
        for period, value in values.items():
            if country == "croatia" and period < "2023-01":
                national = monthly_fields(policy_rows["HR"]).get(period)
                row = policy_rows["HR"]
                selected_value = national
                scope = "country"
                regime = "Croatian national monetary regime before euro adoption"
                policy_regime_id = "croatia_national_pre_euro"
                policy_instrument_id = "bis_cbpol_hr_national_chain"
                series_instance_id = "bis_cbpol_croatia_national"
                regime_start, regime_end = START, "2022-12"
            else:
                row = source_row
                selected_value = value
                scope = "euro_area_common" if use_common else "country"
                regime = "ECB common monetary-policy regime" if use_common else "sovereign monetary-policy regime"
                if use_common:
                    policy_regime_id = "ecb_euro_area_common"
                    policy_instrument_id = "ecb_selected_policy_rate_chain"
                    series_instance_id = "bis_cbpol_euro_area"
                    regime_start, regime_end = ("2023-01" if country == "croatia" else START), None
                elif country == "hungary" and period < "2015-09":
                    policy_regime_id = "hungary_national_policy_chain"
                    policy_instrument_id = "hungary_two_week_deposit_rate"
                    series_instance_id = "bis_cbpol_hungary_two_week_deposit"
                    regime_start, regime_end = START, "2015-08"
                elif country == "hungary":
                    policy_regime_id = "hungary_national_policy_chain"
                    policy_instrument_id = "hungary_three_month_deposit_rate"
                    series_instance_id = "bis_cbpol_hungary_three_month_deposit"
                    regime_start, regime_end = "2015-09", None
                else:
                    policy_regime_id = f"{country}_national_policy_chain"
                    policy_instrument_id = f"bis_cbpol_{geo.lower()}_selected_rate_chain"
                    series_instance_id = f"bis_cbpol_{country}_national"
                    regime_start, regime_end = START, None
            is_shared = policy_regime_id == "ecb_euro_area_common"
            observations.append(observation(
                "policy_rate", country, scope, period, selected_value, "% p.a.", "level",
                "regional_common_driver" if is_shared else "domestic_policy_driver", "BIS / contributing central bank", BIS_FILES["cbpol"],
                "BIS WS_CBPOL", "BIS WS_CBPOL monthly end-of-period", "end_of_month",
                "higher value means a higher selected policy-rate level",
                source_series_id=row.get("Series"), instrument_name=row.get("TITLE", "").strip(),
                instrument_regime=regime, effective_date=period,
                series_break_status="instrument_regime_change_recorded" if row.get("SUPP_INFO_BREAKS") or country == "croatia" else "none_recorded",
                country_monetary_regime=regime, scope_note="ECB common rate is not a country-specific policy decision" if scope == "euro_area_common" else None,
                instrument_regime_metadata=row.get("SUPP_INFO_BREAKS") or None,
                policy_regime_id=policy_regime_id, policy_instrument_id=policy_instrument_id,
                source_series_instance_id=series_instance_id, series_instance_id=series_instance_id,
                regime_start=regime_start, regime_end=regime_end,
                applicability_scope="euro_area" if is_shared else "country",
                applicable_country_ids=sorted(EURO_AREA) if is_shared else [country],
                shared_series=is_shared, independent_cross_section_unit=not is_shared,
                canonical_storage="resolved_country_view" if is_shared else "country_series",
            ))
        if country == "croatia":
            policy_series_manifest.append({
                "country": country, "status": "available", "source_area": "HR through 2022; XM from 2023",
                "scope": "country_then_euro_area_common", "series_id": [policy_rows["HR"].get("Series"), euro_row.get("Series")],
                "source_reference": [policy_rows["HR"].get("SOURCE_REF"), euro_row.get("SOURCE_REF")],
                "instrument_regime_metadata": "Croatian national series before euro adoption; ECB common regime from 2023-01",
            })
        else:
            policy_series_manifest.append({
                "country": country, "status": "available", "source_area": "XM" if use_common else geo,
                "scope": "euro_area_common" if use_common else "country", "series_id": source_row.get("Series"),
                "source_reference": source_row.get("SOURCE_REF"), "instrument_regime_metadata": source_row.get("SUPP_INFO_BREAKS") or None,
            })
    for country in COUNTRIES.values():
        levels = sorted((item for item in observations if item["driver_id"] == "policy_rate" and item["country"] == country and item["transformation"] == "level"), key=lambda item: item["period"])
        by_period = {item["period"]: item for item in levels}
        for item in levels:
            previous_period = shift_month(item["period"], -1)
            previous = by_period.get(previous_period)
            if previous is None and item["period"] == levels[0]["period"]:
                alignment, derivation, reason = "warmup", "warmup", "transformation_warmup"
            elif previous is not None and previous.get("policy_regime_id") != item.get("policy_regime_id"):
                alignment, derivation, reason = "regime_break", "regime_blocked", "policy_regime_transition"
            elif previous is not None and (previous.get("policy_instrument_id") != item.get("policy_instrument_id") or previous.get("series_instance_id") != item.get("series_instance_id")):
                alignment, derivation, reason = "definition_break", "definition_blocked", "policy_instrument_or_source_series_transition"
            elif previous is None or previous["value"] is None or item["value"] is None:
                alignment, derivation, reason = "source_gap", "gap_blocked", "non_consecutive_source_period"
            else:
                alignment, derivation, reason = "aligned", "valid", None
            change = (item["value"] - previous["value"]) * 100 if derivation == "valid" else None
            observations.append(observation(
                "policy_rate", country, item["scope"], item["period"], change, "basis points", "monthly_change_bp",
                item["economic_role"], item["source"], item["source_url"], item["source_dataset"],
                item["definition_version"], item["aggregation_method"],
                "positive means a policy-rate increase in basis points", data_status="computed",
                calculation_method="100 * (policy_rate_t - policy_rate_t-1)",
                source_periods=[previous_period, item["period"]],
                source_observation_ids=[previous["observation_id"], item["observation_id"]] if previous else [item["observation_id"]],
                required_lag_months=1, actual_lag_months=1 if previous else None,
                temporal_alignment_status=alignment, derivation_status=derivation, availability_reason=reason,
                instrument_regime=item.get("instrument_regime"), series_break_status=item.get("series_break_status"),
                policy_regime_id=item.get("policy_regime_id"), policy_instrument_id=item.get("policy_instrument_id"),
                source_series_instance_id=item.get("series_instance_id"), series_instance_id=item.get("series_instance_id"),
                regime_start=item.get("regime_start"), regime_end=item.get("regime_end"),
                applicability_scope=item.get("applicability_scope"), applicable_country_ids=item.get("applicable_country_ids"),
                shared_series=item.get("shared_series"), independent_cross_section_unit=item.get("independent_cross_section_unit"),
            ))
    manifests["policy_rate_acquisition_manifest.json"] = {
        "schema_version": f"policy-rate-acquisition-{SCHEMA_VERSION}", "retrieval_date": TODAY,
        "dataset": "BIS WS_CBPOL", "download_url": BIS_FILES["cbpol"], "frequency": "monthly",
        "file_sha256": sha256(bis_paths["cbpol"][0]), "metadata_version": "current BIS bulk file at retrieval",
        "series": policy_series_manifest,
    }

    # Long-term government yields: Eurostat Maastricht-criterion monthly averages.
    yield_rows, _, yield_url, yield_cache = eurostat("irt_lt_mcby_m", {}, [geo for geo in COUNTRIES if geo != "RS"])
    for row in yield_rows:
        country = COUNTRIES.get(row.get("geo"))
        if country:
            observations.append(observation(
                "long_term_government_yield", country, "country", row["period"], row["value"], "% p.a.", "level",
                "domestic_financial_condition", "Eurostat", "https://ec.europa.eu/eurostat/databrowser/view/irt_lt_mcby_m/default/table",
                "Eurostat irt_lt_mcby_m", "Maastricht criterion bond yield, monthly", "monthly_average",
                "higher value means a higher approximately 10-year government borrowing yield",
                source_series_id=f"irt_lt_mcby_m:M:MCBY:{row.get('geo')}", maturity_definition="approximately 10 years",
            ))
    manifests["interest_rate_acquisition_manifest.json"] = {
        "schema_version": f"interest-rate-acquisition-{SCHEMA_VERSION}", "retrieval_date": TODAY,
        "dataset": "Eurostat irt_lt_mcby_m", "api_query": yield_url, "response_rows": len(yield_rows),
        "response_sha256": sha256(yield_cache), "aggregation_method": "monthly_average",
        "country_status": [{"country": country, "status": "missing_not_substituted" if country == "serbia" else "available",
                            "note": "No comparable official 10-year series found; no alternative maturity substituted." if country == "serbia" else None}
                           for country in COUNTRIES.values()],
    }

    # Bilateral FX from BIS USD cross-rates, with an explicit local-currency-per-EUR orientation.
    with bis_paths["xru"][1].open(encoding="utf-8-sig", newline="") as handle:
        xru_rows = list(csv.DictReader(handle))
    xru_index = {(row["REF_AREA"], row.get("CURRENCY"), row.get("FREQ"), row.get("COLLECTION")): row for row in xru_rows}
    eur_usd = monthly_fields(xru_index[("XM", "EUR", "M", "A")])  # EUR per USD
    common_usd_per_eur = {period: (1 / value if value else None) for period, value in eur_usd.items()}
    add_changes(observations, "eur_usd_common", None, "euro_area", common_usd_per_eur, "USD per EUR",
                "external_common_driver", "BIS", BIS_FILES["xru"], "BIS WS_XRU M:XM:EUR:A",
                "BIS cross-rate, euro-area common series", "monthly_average", "USD per 1 EUR; increase means EUR appreciation",
                series_instance_id="bis_eurusd_common", applicability_scope="european_external_common",
                applicable_country_ids=sorted(COUNTRIES.values()), shared_series=True, independent_cross_section_unit=False)
    fx_manifest_series = []
    for geo, country in COUNTRIES.items():
        if country not in NON_EURO:
            fx_manifest_series.append({"country": country, "series": "eur_usd_common", "scope": "euro_area", "country_specific": False})
            continue
        currency = CURRENCIES[country]
        row = xru_index.get((geo, currency, "M", "A"))
        if not row:
            fx_manifest_series.append({"country": country, "status": "missing"})
            continue
        local_usd = monthly_fields(row)
        cross = {period: (local_usd.get(period) / eur_usd.get(period) if local_usd.get(period) and eur_usd.get(period) else None) for period in sorted(set(local_usd) | set(eur_usd))}
        add_changes(observations, "bilateral_fx_local_per_eur", country, "country", cross, f"{currency} per EUR",
                    "domestic_financial_condition", "BIS (computed cross-rate)", BIS_FILES["xru"],
                    f"BIS WS_XRU {row.get('Series')} and M:XM:EUR:A", "local-per-USD divided by EUR-per-USD",
                    "monthly_average", "local currency units per 1 EUR; increase means local-currency depreciation",
                    calculation_method="(local currency per USD) / (EUR per USD)", numerator_series=row.get("Series"), denominator_series="M:XM:EUR:A",
                    series_instance_id=f"bis_xru_{country}_local_per_eur")
        fx_manifest_series.append({"country": country, "series": row.get("Series"), "currency": currency, "status": "computed_cross_rate"})

    # BIS nominal and real effective exchange rates, broad basket, monthly averages.
    with bis_paths["eer"][1].open(encoding="utf-8-sig", newline="") as handle:
        eer_rows = list(csv.DictReader(handle))
    for geo, country in COUNTRIES.items():
        for eer_type, driver_id in (("N", "nominal_effective_exchange_rate"), ("R", "real_effective_exchange_rate")):
            row = next((item for item in eer_rows if item.get("FREQ") == "M" and item.get("EER_TYPE") == eer_type and item.get("EER_BASKET") == "B" and item.get("REF_AREA") == geo), None)
            if not row:
                continue
            add_changes(observations, driver_id, country, "country", monthly_fields(row), "index",
                        "domestic_financial_condition", "BIS", BIS_FILES["eer"], f"BIS WS_EER {row.get('Series')}",
                        f"BIS broad-basket {'nominal' if eer_type == 'N' else 'real'} EER", "monthly_average",
                        "increase means effective appreciation or strengthening under the BIS convention; it does not by itself imply overvaluation",
                        source_series_id=row.get("Series"), series_instance_id=f"bis_eer_{row.get('Series')}")
    manifests["exchange_rate_acquisition_manifest.json"] = {
        "schema_version": f"exchange-rate-acquisition-{SCHEMA_VERSION}", "retrieval_date": TODAY,
        "datasets": [
            {"dataset": "BIS WS_XRU", "download_url": BIS_FILES["xru"], "file_sha256": sha256(bis_paths["xru"][0])},
            {"dataset": "BIS WS_EER", "download_url": BIS_FILES["eer"], "file_sha256": sha256(bis_paths["eer"][0])},
        ],
        "bilateral_orientation": "local currency units per 1 EUR; increase means local-currency depreciation",
        "effective_fx_orientation": "Under the BIS convention an increase in NEER or REER means effective appreciation/strengthening; it does not by itself imply overvaluation.",
        "effective_fx_methodology_url": "https://data.bis.org/topics/EER",
        "monthly_aggregation": "monthly_average", "series": fx_manifest_series,
        "euro_area_boundary": "EUR/USD is a common external variable and is not represented as a country-specific exchange-rate decision.",
    }

    # Domestic HICP Energy outcome.
    energy_eurostat = []
    for unit, driver_id, result_unit, transformation in (("I15", "hicp_energy_index", "index 2015=100", "level"), ("RCH_A", "hicp_energy_annual_rate", "%", "yoy_rate")):
        rows, _, url, cache = eurostat("prc_hicp_minr", {"coicop18": "NRG", "unit": unit}, list(COUNTRIES))
        for row in rows:
            country = COUNTRIES.get(row.get("geo"))
            if country:
                observations.append(observation(
                    driver_id, country, "country", row["period"], row["value"], result_unit, transformation,
                    "domestic_price_outcome", "Eurostat", "https://ec.europa.eu/eurostat/databrowser/view/prc_hicp_minr/default/table",
                    "Eurostat prc_hicp_minr ECOICOP-2 NRG", f"prc_hicp_minr:NRG:{unit}", "monthly_observation",
                    "domestic consumer energy-price outcome; not an exogenous shock", coicop18="NRG",
                    series_instance_id=f"eurostat_hicp_energy_{unit.lower()}_{country}",
                ))
        energy_eurostat.append({"driver_id": driver_id, "api_query": url, "response_rows": len(rows), "response_sha256": sha256(cache)})

    # Global/European commodity drivers from the World Bank Pink Sheet.
    workbook_path = download(PINK_SHEET_URL, CACHE / "CMO-Historical-Data-Monthly.xlsx")
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheet = workbook["Monthly Prices"]
    sheet_rows = list(sheet.iter_rows(values_only=True))
    names = list(sheet_rows[4])
    units = list(sheet_rows[5])
    commodity_specs = {
        "Crude oil, Brent": ("brent_crude_price_usd", "global"),
        "Natural gas, Europe": ("europe_natural_gas_price_usd", "europe"),
    }
    commodity_manifest = []
    for source_name, (driver_id, scope) in commodity_specs.items():
        column = names.index(source_name) + 1
        values: dict[str, float | None] = {}
        for row in sheet_rows[6:]:
            raw_period = row[0]
            match = re.fullmatch(r"(\d{4})M(\d{2})", str(raw_period or ""))
            if not match:
                continue
            period = f"{match.group(1)}-{match.group(2)}"
            if period < START:
                continue
            raw_value = row[column - 1]
            values[period] = float(raw_value) if isinstance(raw_value, (int, float)) else None
        unit = str(units[column - 1])
        add_changes(observations, driver_id, None, scope, values, unit, "external_common_driver", "World Bank Commodity Price Data (Pink Sheet)",
                    PINK_SHEET_URL, f"World Bank Pink Sheet / Monthly Prices / {source_name}",
                    f"World Bank monthly commodity price: {source_name}", "monthly_average",
                    "USD commodity price level; positive change means price increase", include_yoy=True,
                    series_instance_id="worldbank_brent" if driver_id == "brent_crude_price_usd" else "worldbank_europe_gas",
                    applicability_scope=scope,
                    applicable_country_ids=sorted(COUNTRIES.values()), shared_series=True, independent_cross_section_unit=False)
        commodity_manifest.append({"driver_id": driver_id, "sheet_series_name": source_name, "unit": unit, "rows": len(values)})
    manifests["energy_driver_acquisition_manifest.json"] = {
        "schema_version": f"energy-driver-acquisition-{SCHEMA_VERSION}", "retrieval_date": TODAY,
        "eurostat_hicp_energy": energy_eurostat,
        "pink_sheet": {"source_file": workbook_path.name, "source_url": PINK_SHEET_URL, "sheet": "Monthly Prices",
                       "file_sha256": sha256(workbook_path), "series": commodity_manifest},
        "identification_boundary": "HICP Energy is a domestic outcome; commodity-price movements are external drivers or shock candidates, not identified structural shocks.",
    }

    observations.sort(key=lambda item: item["observation_id"])
    write_json("macro_driver_observations.json", {
        "schema_version": f"macro-driver-observations-{SCHEMA_VERSION}", "generated_at": TODAY,
        "record_count": len(observations), "start_period": START,
        "revision_policy": "latest_revised; first-published vintages are retained only when an official vintage is available",
        "concept_boundary": "Economic role and identification status are orthogonal. A shared observed driver or shock candidate is not an identified shock.",
        "records": observations,
    })

    dictionary = driver_dictionary()
    write_json("macro_driver_dictionary.json", {
        "schema_version": f"macro-driver-dictionary-{SCHEMA_VERSION}", "generated_at": TODAY,
        "record_count": len(dictionary), "records": dictionary,
    })
    coverage = build_coverage(observations)
    write_json("macro_driver_coverage.json", {
        "schema_version": f"macro-driver-coverage-{SCHEMA_VERSION}", "generated_at": TODAY,
        "record_count": len(coverage), "records": coverage,
    })
    for name, payload in manifests.items():
        write_json(name, payload)

    applicability = driver_applicability_registry()
    write_json("driver_applicability_registry.json", {
        "schema_version": f"driver-applicability-registry-{SCHEMA_VERSION}", "generated_at": TODAY,
        "record_count": len(applicability), "records": applicability,
        "boundary": "Applicability does not create statistically independent copies of a shared series.",
    })

    candidates = identified_shock_source_candidates()
    write_json("identified_shock_source_candidates.json", {
        "schema_version": f"identified-shock-source-candidates-{SCHEMA_VERSION}", "generated_at": TODAY,
        "record_count": len(candidates), "records": candidates,
        "boundary": "Candidate registration is not ingestion, validation or identification approval.",
    })
    readiness = v16_identification_readiness()
    write_json("v16_identification_readiness.json", {
        "schema_version": "identification-readiness-v1.6-planning", "generated_at": TODAY,
        "target_release": "v1.6", "all_gates_passed": False,
        "record_count": len(readiness), "records": readiness,
    })

    shock_registry = shock_identification_registry()
    write_json("shock_identification_registry.json", {
        "schema_version": f"shock-identification-registry-{SCHEMA_VERSION}", "generated_at": TODAY,
        "identification_levels": ["observed_driver", "shock_candidate", "external_innovation_proxy", "identified_shock", "blocked"],
        "identified_shock_count": sum(item["identification_status"] == "identified_shock" for item in shock_registry),
        "record_count": len(shock_registry),
        "records": shock_registry,
    })
    lp = lp_readiness(observations, shock_registry)
    write_json("lp_readiness_registry.json", {
        "schema_version": f"lp-readiness-registry-{SCHEMA_VERSION}", "generated_at": TODAY,
        "method_state": "registry_only", "causal_lp_ready_count": sum(item["causal_lp_ready"] for item in lp),
        "record_count": len(lp),
        "readiness_unit": "shock series × outcome × country × sample × controls × horizon",
        "records": lp,
    })
    print(f"Macro drivers: {len(observations)} observations; {len(coverage)} coverage rows; {len(lp)} LP readiness rows.")


def driver_dictionary() -> list[dict]:
    base = [
        ("policy_rate", "政策利率", "Central bank policy rate", "domestic_policy_driver", "country or euro-area common", "% p.a.", "BIS", ["level", "monthly_change_bp"], "Central-bank selected policy-rate level or observed monthly basis-point change; not a monetary-policy shock.", "observed_driver", "Euro-area common rate is not a country-specific policy decision; instrument regimes can change."),
        ("long_term_government_yield", "长期政府债券收益率", "Long-term government bond yield", "domestic_financial_condition", "country", "% p.a.", "Eurostat", ["level"], "Monthly average yield for approximately 10-year central-government bonds.", "observed_driver", "Serbia remains missing because a comparable official series was not established."),
        ("bilateral_fx_local_per_eur", "本币兑欧元汇率", "Bilateral FX: local currency per EUR", "domestic_financial_condition", "country", "local currency per EUR", "BIS computed cross-rate", ["level", "monthly_log_change"], "Increase means local-currency depreciation against EUR.", "observed_driver", "Only non-euro currencies; computed from two BIS USD cross-rates."),
        ("eur_usd_common", "欧元兑美元共同汇率", "EUR/USD common exchange rate", "external_common_driver", "euro_area", "USD per EUR", "BIS", ["level", "monthly_log_change"], "Common euro-area external price, not a national FX decision.", "observed_driver", "Shared regional series, never country-specific."),
        ("nominal_effective_exchange_rate", "名义有效汇率", "Nominal effective exchange rate", "domestic_financial_condition", "country", "index", "BIS", ["level", "monthly_log_change"], "Under the BIS convention, a higher broad-basket NEER index means effective appreciation or strengthening.", "observed_driver", "A higher index does not by itself imply overvaluation; basket composition follows BIS methodology."),
        ("real_effective_exchange_rate", "实际有效汇率", "Real effective exchange rate", "domestic_financial_condition", "country", "index", "BIS", ["level", "monthly_log_change"], "Under the BIS convention, a higher broad-basket REER index means effective appreciation or strengthening.", "observed_driver", "A higher index does not by itself imply overvaluation; price/cost deflator and basket methodology follow BIS."),
        ("hicp_energy_index", "HICP 能源指数", "HICP Energy index", "domestic_price_outcome", "country", "index 2015=100", "Eurostat", ["level"], "Domestic consumer energy-price outcome.", "observed_driver", "Not an external or exogenous energy shock."),
        ("hicp_energy_annual_rate", "HICP 能源同比", "HICP Energy annual rate", "domestic_price_outcome", "country", "%", "Eurostat", ["yoy_rate"], "Domestic consumer energy-price annual rate.", "observed_driver", "Not an external or exogenous energy shock."),
        ("brent_crude_price_usd", "布伦特原油价格", "Brent crude oil price", "external_common_driver", "global", "USD/bbl", "World Bank", ["level", "monthly_log_change", "12m_log_change"], "Global observed energy-price movement.", "observed_driver", "Monthly change is registered separately as a shock candidate; neither level nor change is an identified supply or geopolitical shock."),
        ("europe_natural_gas_price_usd", "欧洲天然气价格", "European natural gas price", "external_common_driver", "europe", "USD/mmbtu", "World Bank", ["level", "monthly_log_change", "12m_log_change"], "European observed energy-price movement.", "observed_driver", "Monthly change is registered separately as a shock candidate; neither level nor change is an identified supply or geopolitical shock."),
    ]
    return [{"driver_id": item[0], "name_zh": item[1], "name_en": item[2],
             "economic_role": item[3], "role": item[3], "identification_default": item[9],
             "identification_status": item[9], "frequency": "monthly",
             "series_scope_type": item[4], "scope": item[4], "shared_series": item[4] in {"global", "europe", "euro_area", "country or euro-area common"},
             "timing_convention": "monthly end-of-period" if item[0] == "policy_rate" else "monthly observation or monthly average as documented",
             "regime_sensitive": item[0] == "policy_rate",
             "unit": item[5], "source": item[6], "transformation_options": item[7],
             "economic_interpretation": item[8], "limitations": item[10]}
            for item in base]


def build_coverage(records: list[dict]) -> list[dict]:
    grouped: dict[tuple, list[dict]] = defaultdict(list)
    for record in records:
        grouped[(record["driver_id"], record["country"], record["scope"], record["transformation"])].append(record)
    result = []
    for (driver_id, country, scope, transformation), rows in sorted(grouped.items(), key=lambda item: str(item[0])):
        periods = sorted(row["period"] for row in rows)
        expected = month_sequence(periods[0], periods[-1]) if periods else []
        values = {row["period"]: row["value"] for row in rows}
        level_rows = [row for row in records if row["driver_id"] == driver_id and row["country"] == country and row["scope"] == scope and row["transformation"] == "level"]
        source_rows = level_rows or rows
        source_values = {row["period"]: row["value"] for row in source_rows}
        source_periods = sorted(source_values)
        source_expected = month_sequence(source_periods[0], source_periods[-1]) if source_periods else []
        source_missing = [period for period in source_expected if source_values.get(period) is None]
        warmup = [row["period"] for row in rows if row.get("derivation_status") == "warmup"]
        gaps = [row["period"] for row in rows if row.get("derivation_status") == "gap_blocked"]
        regime_excluded = [row["period"] for row in rows if row.get("derivation_status") == "regime_blocked"]
        definition_excluded = [row["period"] for row in rows if row.get("derivation_status") == "definition_blocked"]
        transformed_observations = sum(value is not None for value in values.values())
        missing = sorted(set(source_missing + gaps + regime_excluded + definition_excluded))
        transformed_eligible_periods = max(0, len(expected) - len(warmup))
        raw_source_coverage = (sum(value is not None for value in source_values.values()) / len(source_expected)) if source_expected else 0
        transformed_usable_coverage = (transformed_observations / transformed_eligible_periods) if transformed_eligible_periods else 0
        breaks = sorted(set(row.get("series_break_status") for row in rows if row.get("series_break_status") not in (None, "none_recorded")))
        result.append({
            "driver_id": driver_id, "country": country, "scope": scope, "transformation": transformation,
            "start_period": periods[0] if periods else None, "end_period": periods[-1] if periods else None,
            "expected_periods": len(expected), "observations": transformed_observations,
            "missing_periods": missing, "definition_status": "documented",
            "source_expected_periods": len(source_expected), "source_observations": sum(value is not None for value in source_values.values()),
            "source_missing_periods": source_missing,
            "transformed_expected_periods": len(expected), "transformed_observations": transformed_observations,
            "transformation_warmup_periods": warmup, "transformation_gap_periods": gaps,
            "regime_excluded_periods": regime_excluded, "definition_excluded_periods": definition_excluded,
            "raw_source_coverage": round(raw_source_coverage, 6),
            "transformed_eligible_periods": transformed_eligible_periods,
            "transformed_usable_coverage": round(transformed_usable_coverage, 6),
            "warmup_count": len(warmup), "source_gap_count": len(gaps),
            "regime_exclusion_count": len(regime_excluded), "definition_exclusion_count": len(definition_excluded),
            "series_break_status": "none_recorded" if not breaks else ";".join(breaks),
            "analysis_eligible": bool(expected) and transformed_usable_coverage >= 0.9,
            "eligibility_basis": "documented definition + exact temporal alignment + regime/definition exclusions removed + usable coverage >= 90%",
        })
    return result


def driver_applicability_registry() -> list[dict]:
    records = []
    for country in sorted(EURO_AREA - {"croatia"}):
        records.append({"applicability_id": f"bis_cbpol_euro_area:{country}:2015", "series_instance_id": "bis_cbpol_euro_area",
                        "driver_id": "policy_rate", "country_id": country, "start_period": START, "end_period": None,
                        "applicability_scope": "euro_area", "shared_series": True, "independent_cross_section_unit": False})
    records.append({"applicability_id": "bis_cbpol_euro_area:croatia:2023", "series_instance_id": "bis_cbpol_euro_area",
                    "driver_id": "policy_rate", "country_id": "croatia", "start_period": "2023-01", "end_period": None,
                    "applicability_scope": "euro_area", "shared_series": True, "independent_cross_section_unit": False,
                    "note": "Croatia entered the euro area in 2023-01; earlier policy observations remain in the national regime."})
    for series_id, driver_id, scope in (
        ("bis_eurusd_common", "eur_usd_common", "european_external_common"),
        ("worldbank_brent", "brent_crude_price_usd", "global"),
        ("worldbank_europe_gas", "europe_natural_gas_price_usd", "europe"),
    ):
        for country in sorted(COUNTRIES.values()):
            records.append({"applicability_id": f"{series_id}:{country}", "series_instance_id": series_id,
                            "driver_id": driver_id, "country_id": country, "start_period": START, "end_period": None,
                            "applicability_scope": scope, "shared_series": True, "independent_cross_section_unit": False})
    return records


def identified_shock_source_candidates() -> list[dict]:
    return [
        {"candidate_id": "ecb_ea_mpd", "name": "Euro Area Monetary Policy Event-Study Database (EA-MPD)",
         "source_institution": "European Central Bank", "source_url": "https://www.ecb.europa.eu/pub/pdf/annex/Dataset_EA-MPD.xlsx",
         "methodology_url": "https://www.ecb.europa.eu/pub/pdf/scpwps/ecb.wp2281~3303fd281b.en.pdf",
         "authors": ["Carlo Altavilla", "Luca Brugnolini", "Refet S. Gürkaynak", "Roberto Motto", "Giuseppe Ragusa"],
         "dataset_type": "official ECB research dataset; event-level high-frequency/intraday surprise information",
         "coverage_scope": "euro_area_policy_announcements", "identification_status": "candidate_for_identification_review",
         "required_audits": ["variable-definition", "information-effect", "event-window", "aggregation policy", "coverage", "license", "vintage", "country applicability"],
         "croatia_interpretation": "Euro-area surprise exposure can apply only after euro adoption unless a separate research design justifies earlier external spillover use."},
        {"candidate_id": "ecb_ea_empd", "name": "Euro Area Extended Monetary Policy Event-Study Database (EA-EMPD)",
         "source_institution": "European Central Bank", "source_url": "https://www.ecb.europa.eu/pub/pdf/scpwps/ecb.wp3157~f5789c7b8a.da.pdf",
         "methodology_url": "https://www.ecb.europa.eu/pub/pdf/scpwps/ecb.wp3157~f5789c7b8a.da.pdf",
         "authors": ["Carlo Altavilla", "Refet S. Gürkaynak", "Luc Laeven", "Thilo Kind"],
         "dataset_type": "official ECB research annex; official decisions and policymaker speeches",
         "official_annex_status": "listed by ECB; exact downloadable annex asset remains subject to acquisition and checksum audit",
         "coverage_scope": "euro_area_decisions_and_policy_speeches", "identification_status": "candidate_for_identification_review",
         "required_audits": ["annex download", "variable-definition", "information-effect", "event-window", "aggregation policy", "coverage", "license", "revision policy", "country applicability"],
         "croatia_interpretation": "Country treatment must distinguish euro-area membership from external spillover exposure and preserve event-date applicability."},
    ]


def v16_identification_readiness() -> list[dict]:
    return [
        {"gate_id": "ea_mpd_acquisition", "status": "not_started", "required_evidence": "official EA-MPD asset, immutable checksum and redistribution review"},
        {"gate_id": "ea_empd_acquisition", "status": "not_started", "required_evidence": "official EA-EMPD annex asset, immutable checksum and redistribution review"},
        {"gate_id": "shock_field_definitions", "status": "pending", "required_evidence": "registered factors, units, signs and information-set interpretation"},
        {"gate_id": "announcement_timestamps", "status": "pending", "required_evidence": "decision/speech timestamps, market windows and timezone handling"},
        {"gate_id": "event_coverage", "status": "pending", "required_evidence": "event universe, omissions, revisions and vintage audit"},
        {"gate_id": "monthly_aggregation", "status": "pending", "required_evidence": "pre-registered event-to-month aggregation without look-ahead"},
        {"gate_id": "information_effect_handling", "status": "pending", "required_evidence": "separate policy and information components or an explicit limitation"},
        {"gate_id": "euro_area_applicability", "status": "pending", "required_evidence": "dated monetary-union applicability including Croatia from 2023-01"},
        {"gate_id": "non_euro_spillover_interpretation", "status": "pending", "required_evidence": "external ECB spillover wording; never domestic monetary-policy shock"},
        {"gate_id": "lp_estimator_validation", "status": "blocked", "required_evidence": "separate LP specification, inference and diagnostic validation"},
    ]


def shock_identification_registry() -> list[dict]:
    records = []
    for driver_id in ["policy_rate", "long_term_government_yield", "bilateral_fx_local_per_eur", "eur_usd_common", "nominal_effective_exchange_rate", "real_effective_exchange_rate", "hicp_energy_index", "hicp_energy_annual_rate"]:
        records.append({"shock_id": f"{driver_id}_movement", "driver_id": driver_id, "transformation": "level_or_change",
                        "identification_status": "observed_driver", "identification_strategy": None,
                        "causal_use_allowed": False, "boundary": "Observed movement is not an identified shock."})
    for driver_id in ["brent_crude_price_usd", "europe_natural_gas_price_usd"]:
        records.append({"shock_id": f"{driver_id}_monthly_change", "driver_id": driver_id, "transformation": "monthly_log_change",
                        "identification_status": "shock_candidate", "identification_strategy": None,
                        "causal_use_allowed": False, "boundary": "Observed commodity-price change is not an identified structural energy shock."})
    records.extend([
        {"shock_id": "monetary_policy_surprise", "driver_id": "policy_rate", "transformation": "announcement_surprise",
         "identification_status": "blocked", "identification_strategy": "requires a validated high-frequency announcement surprise or external instrument",
         "causal_use_allowed": False, "boundary": "A policy-rate change is not a monetary-policy shock."},
        {"shock_id": "external_energy_price_innovation", "driver_id": "europe_natural_gas_price_usd", "transformation": "innovation",
         "identification_status": "blocked", "identification_strategy": "innovation construction not yet registered or validated",
         "causal_use_allowed": False, "boundary": "Price innovation cannot be called a supply or geopolitical shock without identification."},
    ])
    return records


def lp_readiness(records: list[dict], shock_registry: list[dict]) -> list[dict]:
    hf_path = ROOT / "src" / "data" / "high-frequency" / "high_frequency_observations.json"
    hf = json.loads(hf_path.read_text(encoding="utf-8"))["records"]
    hf_group: dict[tuple[str, str], list[dict]] = defaultdict(list)
    for row in hf:
        if row.get("value") is not None:
            hf_group[(row["country"], row["indicator"])].append(row)
    policy_periods: dict[str, set[str]] = defaultdict(set)
    for row in records:
        if row["driver_id"] == "policy_rate" and row["transformation"] == "level" and row["value"] is not None and row["country"]:
            policy_periods[row["country"]].add(row["period"])
    result = []
    candidates = [item for item in shock_registry if item["identification_status"] == "shock_candidate"]
    for candidate in candidates:
        candidate_rows = [row for row in records if row["driver_id"] == candidate["driver_id"] and row["transformation"] == candidate["transformation"]]
        shock_periods = {row["period"] for row in candidate_rows if row["value"] is not None and row.get("derivation_status") == "valid" and row.get("temporal_alignment_status") == "aligned"}
        warmup_periods = sorted({row["period"] for row in candidate_rows if row.get("derivation_status") == "warmup"})
        excluded_periods = sorted({row["period"] for row in candidate_rows if row.get("derivation_status") in {"gap_blocked", "regime_blocked", "definition_blocked"}})
        shock_scope = candidate_rows[0].get("applicability_scope", candidate_rows[0].get("scope")) if candidate_rows else None
        for country in COUNTRIES.values():
            for outcome in ("hicp_annual_rate", "industrial_production_index"):
                outcome_periods = {row["period"] for row in hf_group.get((country, outcome), [])}
                common = sorted(shock_periods & outcome_periods & policy_periods.get(country, set()))
                result.append({
                    "readiness_id": f"lp:{candidate['shock_id']}:{outcome}:{country}:h0-24",
                    "shock_series": candidate["shock_id"], "outcome": outcome, "country": country,
                    "sample_start": common[0] if common else None, "sample_end": common[-1] if common else None,
                    "controls": ["policy_rate"], "horizon": list(range(25)), "monthly_observations": len(common),
                    "shock_scope": shock_scope, "shared_series_instance_id": candidate_rows[0].get("series_instance_id") if candidate_rows else None,
                    "transformation_warmup_periods": warmup_periods, "temporally_or_definition_excluded_periods": excluded_periods,
                    "minimum_observations_met": len(common) >= 96, "shock_definition_complete": False,
                    "outcome_coverage_complete": len(outcome_periods) >= 96, "controls_complete": len(policy_periods.get(country, set())) >= 96,
                    "unresolved_definition_break": False, "identification_status": candidate["identification_status"],
                    "data_ready": len(common) >= 96 and not excluded_periods, "causal_lp_ready": False,
                    "method_state": "registry_only",
                    "blockers": ["shock identification is not identified_shock", "Local Projections estimator is not activated in v1.51"],
                })
    return result


if __name__ == "__main__":
    build()
